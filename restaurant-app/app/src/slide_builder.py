"""
Slide Builder: constructs PowerPoint presentations (.pptx) adhering to waCRM design guidelines.
"""
import os
from typing import List, Dict, Any
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

from .config import (
    SLIDE_WIDTH_PX,
    SLIDE_HEIGHT_PX,
    px_to_in,
    COLOR_BLUE,
    COLOR_GRAY,
    COLOR_GREEN,
    FONT_FAMILY,
    PC_POS_X,
    PC_POS_Y,
    PC_WIDTH,
    PC_HEIGHT,
    TEXT_COL_X,
    KICKER_Y,
    HEADING_Y,
    BODY_Y,
    BULLETS_START_Y,
    BULLET_SPACING_Y,
    CHECK_ICON_SIZE,
    ENDING_Y
)

def prepare_background(bg_trans_path: str, output_path: str) -> str:
    """
    Composites the transparent background template over a clean white 1920x1080 canvas.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    if os.path.exists(output_path):
        return output_path

    white_bg = Image.new("RGBA", (SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX), (255, 255, 255, 255))
    trans_bg = Image.open(bg_trans_path).convert("RGBA")
    composite = Image.alpha_composite(white_bg, trans_bg)
    composite.save(output_path, "PNG")
    return output_path

def add_wacrm_slide(
    prs: Presentation,
    content: Dict[str, Any],
    bg_image_path: str,
    pc_mockup_path: str,
    check_icon_path: str
):
    """
    Adds a fully styled waCRM slide to the presentation.
    """
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Background image
    slide.shapes.add_picture(bg_image_path, 0, 0, prs.slide_width, prs.slide_height)

    # 2. PC Mockup on the left
    slide.shapes.add_picture(
        pc_mockup_path,
        px_to_in(PC_POS_X),
        px_to_in(PC_POS_Y),
        px_to_in(PC_WIDTH),
        px_to_in(PC_HEIGHT)
    )

    # 3. Kicker Text (e.g. SETTINGS & SETUP)
    tx_kicker = slide.shapes.add_textbox(
        px_to_in(TEXT_COL_X),
        px_to_in(KICKER_Y),
        Inches(4.8),
        Inches(0.4)
    )
    tf_k = tx_kicker.text_frame
    tf_k.word_wrap = True
    tf_k.margin_left = tf_k.margin_right = tf_k.margin_top = tf_k.margin_bottom = 0
    p_k = tf_k.paragraphs[0]
    p_k.text = content.get("kicker", "").upper()
    p_k.font.name = FONT_FAMILY
    p_k.font.size = Pt(18)
    p_k.font.bold = True
    p_k.font.color.rgb = COLOR_BLUE

    # 4. Heading (e.g. Connect Whatsapp)
    tx_heading = slide.shapes.add_textbox(
        px_to_in(TEXT_COL_X),
        px_to_in(HEADING_Y),
        Inches(5.0),
        Inches(0.75)
    )
    tf_h = tx_heading.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = tf_h.margin_right = tf_h.margin_top = tf_h.margin_bottom = 0
    p_h = tf_h.paragraphs[0]
    p_h.text = content.get("heading", "")
    p_h.font.name = FONT_FAMILY
    p_h.font.size = Pt(40)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_GRAY

    # 5. Body description text
    tx_body = slide.shapes.add_textbox(
        px_to_in(TEXT_COL_X),
        px_to_in(BODY_Y),
        Inches(4.8),
        Inches(0.65)
    )
    tf_b = tx_body.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
    p_b = tf_b.paragraphs[0]
    p_b.text = content.get("body", "")
    p_b.font.name = FONT_FAMILY
    p_b.font.size = Pt(15)
    p_b.font.color.rgb = COLOR_GRAY

    # 6. Bullet Points with checkmark icon
    bullets = content.get("bullets", [])
    y_current = BULLETS_START_Y
    for b in bullets:
        slide.shapes.add_picture(
            check_icon_path,
            px_to_in(TEXT_COL_X),
            px_to_in(y_current),
            px_to_in(CHECK_ICON_SIZE),
            px_to_in(CHECK_ICON_SIZE)
        )
        tx_b = slide.shapes.add_textbox(
            px_to_in(TEXT_COL_X + CHECK_ICON_SIZE + 12),
            px_to_in(y_current - 2),
            Inches(4.5),
            Inches(0.35)
        )
        tf_bl = tx_b.text_frame
        tf_bl.word_wrap = True
        tf_bl.margin_left = tf_bl.margin_right = tf_bl.margin_top = tf_bl.margin_bottom = 0
        p = tf_bl.paragraphs[0]
        p.text = b
        p.font.name = FONT_FAMILY
        p.font.size = Pt(15)
        p.font.color.rgb = COLOR_GRAY

        y_current += BULLET_SPACING_Y

    # 7. Ending Sentence
    ending_y = max(y_current + 15, ENDING_Y)
    tx_ending = slide.shapes.add_textbox(
        px_to_in(TEXT_COL_X),
        px_to_in(ending_y),
        Inches(4.8),
        Inches(0.65)
    )
    tf_e = tx_ending.text_frame
    tf_e.word_wrap = True
    tf_e.margin_left = tf_e.margin_right = tf_e.margin_top = tf_e.margin_bottom = 0
    p_e = tf_e.paragraphs[0]
    p_e.text = content.get("ending", "")
    p_e.font.name = FONT_FAMILY
    p_e.font.size = Pt(15)
    p_e.font.color.rgb = COLOR_GRAY

    return slide
